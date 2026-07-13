"""Catalog XLSX import pipeline (CAT-003).

Feeds workbook rows through the SAME validation core as the CSV
pipeline (CAT-002) — identical issue reporting, formula flagging, and
bounds — with the workbook-specific hazards handled at the reading
layer:

- **macros are never executed** — openpyxl has no VBA engine; workbooks
  are opened read-only without ``keep_vba``, so a macro-carrying .xlsm
  is read as inert data;
- **formulas are never evaluated** — the workbook is opened with
  ``data_only=True``: formula cells yield the CACHED value the authoring
  application last computed, or empty when none exists (reported as a
  file-level warning is not possible per-cell with cached reads, so the
  guarantee is simpler: nothing in this platform ever evaluates a
  formula). String values that LOOK like formulas get the same warning
  and export defusing as CSV cells;
- **sheet selection is explicit** — an unknown sheet name fails loudly
  listing what the workbook actually contains; omitting the name uses
  the first sheet;
- cells are coerced conservatively: dates become ISO strings, integral
  floats lose their ``.0``, booleans become ``true``/``false``.
"""

import io
import zipfile
from datetime import date, datetime
from typing import Any

from soa_db.catalog_import import (
    MAX_ROWS,
    CatalogImportError,
    ColumnMapping,
    ParseResult,
    validate_import_rows,
)


def _coerce(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == value.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def parse_catalog_xlsx(
    data: bytes, mapping: ColumnMapping, *, sheet: str | None = None
) -> ParseResult:
    """Parse and validate an XLSX/XLSM upload. See module docstring."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,  # cached values only; formulas are never evaluated
        )
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError):
        raise CatalogImportError("the file could not be read as an XLSX workbook") from None

    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise CatalogImportError(
                    f"the workbook has no sheet named {sheet!r} "
                    f"(sheets: {', '.join(workbook.sheetnames)})"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]

        rows_iter = worksheet.iter_rows(values_only=True)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            raise CatalogImportError("the sheet is empty") from None
        header = [_coerce(cell) for cell in header_cells]

        def rows() -> "list[dict[str, str]]":
            materialized: list[dict[str, str]] = []
            for index, raw in enumerate(rows_iter):
                if index >= MAX_ROWS:
                    raise CatalogImportError(f"the file exceeds the {MAX_ROWS} row budget")
                values = [_coerce(cell) for cell in raw]
                if not any(values):
                    continue  # fully empty spreadsheet rows are not data
                materialized.append(dict(zip(header, values, strict=False)))
            return materialized

        return validate_import_rows(
            rows(),
            header,
            mapping,
            warnings=[],
            encoding="xlsx",
            delimiter="",
        )
    finally:
        workbook.close()


__all__ = ["parse_catalog_xlsx"]
